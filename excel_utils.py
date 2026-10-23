from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook

import database

def export_products(file_path):
    """ 将数据库中的商品导出围殴 Excel"""
    
    # 读取全部商品
    products = database.get_all_products()
    
    # 创建Excel工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = "商品资料"
    
    # 写入表头
    headers = ["ID", "商品货号", "商品名称", "尺码"]
    ws.append(headers)
    
    # 写入商品数据
    for product in products:
        ws.append(product)
        
    # 设置表头样式
    for cell in ws[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )
        cell.fill = PatternFill(
            "solid",
            fgColor="4472C4"
        )
        cell.alignment = Alignment(
            horizontal = "center"
        )
        
    # 设置列宽
    widths = [10, 22, 35, 15]
    
    for index, width in enumerate(widths, start=1):
        column_letter = get_column_letter(index)
        ws.column_dimensions[column_letter].width=width
        
    # 冻结第一行
    ws.freeze_panes = "A2"
    
    # 设置自动筛选
    ws.auto_filter.ref = ws.dimensions
    
    # 保存Excel文件
    wb.save(file_path)
    
    return len(products)
    
def preview_import_products(file_path):
    """ 读取并检查 Excel 商品数据，不写入数据库"""
    
    wb = load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )
    
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        
        # 读取第一行表头
        headers = next(rows, None)
        
        if headers is None:
            raise ValueError("Excel 文件为空。")
            
        expected = [
            "ID",
            "商品货号",
            "商品名称",
            "尺码"
        ]
        
        actual = [
            str(value).strip() if value is not None else ""
            for value in headers[:4]
        ]
        
        if actual != expected:
            raise ValueError(
                "Excel 表头不正确.\n"
                f"需要： {expected}\n"
                f"实际： {actual}"
            )
            
        products = []
        errors = []
        
        for row_number, row in enumerate(rows, start=2):
            
            # 跳过完全空白的行
            if not any(
                value is not None and str(value).strip()
                for value in row
            ):
                continue
                
            code = row[1] if len(row) > 1 else None
            name = row[2] if len(row) > 2 else None
            size = row[3] if len(row) > 3 else None
            
            code = str(code).strip() if code is not None else ""
            name = str(name).strip() if name is not None else ""
            size = str(size).strip() if size is not None else ""
            
            if not code or not name:
                errors.append(
                    f"第 {row_number} 行：商品货号或商品名称为空"
                )
                continue
                
            
            products.append((code, name, size))
            
        if errors:
            raise ValueError(
                "发现数据错误: \n"
                +"\n".join(errors[:10])
                +(
                    f"\n还有{len(errors)-10}条错误"
                    if len(errors) > 10 else ""
                )
            )
            
        return products
        
    finally:
        wb.close()